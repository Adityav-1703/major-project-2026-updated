#!/usr/bin/env python3
"""
AyurAuth B2C presentation builder (matches the live website).
Slide: 13.333 in x 7.5 in | Font: Times New Roman
Run from this folder:  python pptcode.py
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image as PILImage

# =================================================================
#  SLIDE CONTENT  —  matches the live B2C AyurAuth website
# =================================================================
SLIDES = [
    # ------------------------------------------------------------------
    # SLIDE 1 — TITLE
    # ------------------------------------------------------------------
    {
        "title": "AyurAuth — B2C Herb Authenticity: Scan, Shop, and Trust What You Buy",
        "image": None,
        "bullets": [
            "Under the guidance of: Prof. Jayesh Rane and Vijen Gala",
            "Aditya Rajesh Srinivas — 16014022115",
            "Shreya Vikas Jhamb   — 16014022121",
            "",
            "Product today: a shopper-first (B2C) web app — not a farmer-only portal.",
            "",
            "Tech Stack (as implemented):",
            "  React 18 + TypeScript + Vite + Tailwind CSS + Framer Motion",
            "  Node.js + Express + JWT auth (consumer vs grower roles)",
            "  Local JSON store (fallback) + optional MongoDB Atlas",
            "  CNN — MobileNetV2 (herbal_auth_improved_final.h5), Flask on :5001",
            "  5 classes: Amla, Guava, Hibiscus, Neem, Tulsi  |  81.15% val. accuracy",
            "  QR verify (jsQR) + Shop catalog + My checks (consumer dashboard)",
            "  Shopper chatbot (rule engine + optional LLM)  |  10 languages",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 2 — OUTLINE
    # ------------------------------------------------------------------
    {
        "title": "Outline",
        "image": None,
        "bullets": [
            "1.  Introduction — B2C authenticity for shoppers",
            "2.  Problem Statement",
            "3.  Literature Survey (10 papers)",
            "4.  Literature Review Summary",
            "5.  Objectives of the Project",
            "6.  System Architecture — Block Diagram",
            "7.  Shopper Workflow — Scan, Shop, My checks",
            "8.  Methodology & Proposed Approach",
            "9.  CNN Model — MobileNetV2 (5 classes, measured scores)",
            "10. AI Chatbot — shopper assistant",
            "11. Backend — Express + JWT + local/Mongo storage",
            "12. Work Done (live website)",
            "13. Website screens (routes & features)",
            "14. Conclusion & QnA",
            "15. References",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 3 — INTRODUCTION
    # ------------------------------------------------------------------
    {
        "title": "Introduction",
        "image": None,
        "bullets": [
            "Ayurveda is one of the world's oldest holistic healing systems. Counterfeit and mislabelled herbs put buyers at risk.",
            "AyurAuth is implemented as a B2C web product: the primary user is the shopper holding a pack, not only the farmer.",
            "",
            "What a buyer can do on the live site (port 5173):",
            "  • Scan a pack — camera QR, image upload, or batch ID  (/verify)",
            "  • Browse verified listings with origin, AI score, indicative price  (/shop)",
            "  • Create a free shopper account (default role: consumer)",
            "  • Save scan history under My checks  (/consumer)",
            "  • Ask the shopper chatbot how scanning and AI scores work",
            "",
            "Growers remain supported as a secondary path: Sell / list a batch (/upload) + grower dashboard.",
            "Authenticity is image-based CNN support — it does not replace laboratory testing.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 4 — PROBLEM STATEMENT
    # ------------------------------------------------------------------
    {
        "title": "Problem Statement",
        "image": None,
        "bullets": [
            '"A Blockchain-driven system for authenticity, sustainability and export readiness of ayurvedic herbs"',
            "",
            "Buyers cannot tell a genuine Ayurvedic pack from a fake or mislabelled one at the shelf.",
            "",
            "Problems the current website actually attacks:",
            "1. No instant pack check  —  shoppers now scan QR / batch ID and see AI + origin.",
            "2. No public catalog of verified batches  —  /shop lists AI-verified herbs.",
            "3. Farmer-only UX  —  signup defaults to consumer; landing CTAs are Scan / Shop / Sign up.",
            "4. No saved proof of checks  —  My checks stores personal scan history.",
            "5. Weak visual authenticity  —  MobileNetV2 CNN on 5 leaf classes with a 70% threshold.",
            "6. Opaque records  —  each listing carries a tamper-evident hash + QR payload.",
            "",
            "Out of current product scope (research goals, not live checkout): full Hyperledger network,",
            "EU FMD / DSCSA automation, payments/cart, and live sustainability sensors.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 5 — LITERATURE SURVEY (Part 1)
    # ------------------------------------------------------------------
    {
        "title": "Literature Survey (References 1 – 5)",
        "image": None,
        "bullets": [
            "Ref 1 — Kunjachan et al. (2023): CNN comparison (VGG-16, ResNet-50, EfficientNet-V2M) on 1,835 images / 30 Ayurvedic species.",
            "  Strength: EfficientNet-V2M achieved 98.24% accuracy.  Limitation: Sensitive to lighting and background.",
            "",
            "Ref 2 — Kaushik et al. (2025): AI-enabled blockchain + QR for AYUSH drug supply chains (analysis of e-Aushadhi, e-Charak).",
            "  Strength: Comprehensive regulatory roadmap.  Limitation: Largely conceptual, requires rural infrastructure.",
            "",
            "Ref 3 — Uniyal et al.: Hyperledger Fabric + Sequential CNN on 20-class medicinal plant dataset.",
            "  Strength: Real-time tamper-proof tracking.  Limitation: Lower classification accuracy than standalone DL models.",
            "",
            "Ref 4 — Liu et al. (2021): Segmented Encryption — Blockchain + Smart Contracts + RFID/QR for herbal medicine.",
            "  Strength: Seed-to-shelf monitoring with high transparency.  Limitation: High cost for small-scale farmers.",
            "",
            "Ref 5 — Rout (2025): MedicLedger — Ethereum blockchain for Ayurvedic product authentication.",
            "  Strength: Eliminates intermediaries, ensures fair pricing.  Limitation: Gas fees and scalability challenges.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 6 — LITERATURE SURVEY (Part 2)
    # ------------------------------------------------------------------
    {
        "title": "Literature Survey (References 6 – 10)",
        "image": None,
        "bullets": [
            "Ref 6 — Balaji & Parani (2022): DNA Barcoding (rbcL marker) on 107 powder samples / 65 species.",
            "  Strength: Accurate even in processed powder.  Limitation: Requires expert laboratory setup.",
            "",
            "Ref 7 — Liu et al. (2021): Dual Blockchain (Private + Consortium) + IPFS for rare Chinese herbs.",
            "  Strength: Efficient off-chain storage.  Limitation: Increased technical complexity.",
            "",
            "Ref 8 — Kaushal et al. (2023): Hyperledger Fabric benchmarking with simulated multi-org transactions.",
            "  Strength: High throughput, low latency, strong privacy.  Limitation: Advanced expertise required.",
            "",
            "Ref 9 — Reddy et al. (2024): AI + Predictive Analytics + Blockchain for pharmaceutical supply chain.",
            "  Strength: Accurate demand forecasting.  Limitation: Requires large volumes of high-quality historical data.",
            "",
            "Ref 10 — Yik et al. (2021): HerBChain — Consortium blockchain (Openchain) + QR for herbal quality assurance.",
            "  Strength: Prevents manipulation of quality records.  Limitation: Resistance to adoption in traditional industry.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 7 — LITERATURE REVIEW SUMMARY
    # ------------------------------------------------------------------
    {
        "title": "Literature Review — Summary",
        "image": None,
        "bullets": [
            "While individual technologies exist, no single system addresses all challenges together.",
            "",
            "Existing Systems:",
            "  • Blockchain traceability systems (Hyperledger, Ethereum)",
            "  • CNN models achieving up to 98% accuracy on herb images",
            "  • Government portals — e-Aushadhi, e-Charak",
            "  • QR-based supply chain monitoring",
            "",
            "Gaps Identified across all 10 papers:",
            "  • No unified end-to-end system covering the full supply chain",
            "  • No export compliance automation (EU FMD / US DSCSA)",
            "  • Limited scalability and authentication in existing platforms",
            "  • Weak or absent real-time sustainability tracking",
            "",
            "Conclusion: Literature still lacks a unified buyer-facing authenticity app. AyurAuth implements that B2C layer now.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 8 — OBJECTIVES
    # ------------------------------------------------------------------
    {
        "title": "Objectives of the Project",
        "image": None,
        "bullets": [
            "1. Give shoppers a one-tap pack check: Scan QR, upload QR image, or enter batch ID (/verify).",
            "2. Publish a B2C catalog of AI-verified herbs with origin, confidence, and indicative pack price (/shop).",
            "3. Default new accounts to consumer; growers keep /upload and /dashboard as a secondary path.",
            "4. Save personal scan history in My checks (/consumer) after login (JWT).",
            "5. Run MobileNetV2 CNN (224×224) on Amla, Guava, Hibiscus, Neem, Tulsi; verify if confidence ≥ 0.70.",
            "6. Attach a tamper-evident hash + ayurauth://batch/{uuid} QR to every listed batch.",
            "7. Ship a shopper chatbot (rule engine + optional LLM) that never shares contact numbers.",
            "8. Multilingual React SPA (10 languages) with dark/light theme; Flask ML service on port 5001.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 9 — SYSTEM ARCHITECTURE BLOCK DIAGRAM (text description)
    # ------------------------------------------------------------------
    {
        "title": "System Architecture — Block Diagram",
        "image": None,
        "bullets": [
            "Layer 1 — SHOPPER (B2C) + GROWER:",
            "  Shopper: Scan pack | Shop listings | My checks | Signup/Login",
            "  Grower (secondary): List a batch | Grower dashboard",
            "              ↓  browser  http://localhost:5173",
            "",
            "Layer 2 — FRONTEND  (React 18 + Vite + Tailwind):",
            "  /  /shop  /verify  /consumer  /upload  /dashboard  /login  /signup  /chat",
            "  Language switcher  |  Theme toggle  |  Floating AIChatbot",
            "              ↓  VITE_API_URL → http://localhost:5000",
            "",
            "Layer 3 — BACKEND  (Express :5000, JWT):",
            "  /api/auth  |  /api/herbs  |  /api/verify  |  /api/chat  |  /api/predict  |  /api/health",
            "  Storage: USE_LOCAL_DB → backend/data/local-db.json  (Atlas optional)",
            "              ↓",
            "",
            "Layer 4 — AI  (Flask :5001):",
            "  herbal_auth_improved_final.h5  |  224×224  |  softmax 5 classes  |  threshold 0.70",
            "",
            "Layer 5 — TRUST RECORD:",
            "  UUID batchId  |  fakeBlockchainHash (hex)  |  QR payload ayurauth://batch/{id}",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 10 — SYSTEM WORKFLOW
    # ------------------------------------------------------------------
    {
        "title": "Shopper Workflow — How the Website Works Now",
        "image": None,
        "bullets": [
            "B2C PATH (primary):",
            "  1. Land on home — CTAs: Scan a pack, Browse verified herbs, Sign up",
            "  2. Shop (/shop) — AI-verified cards: origin, AI %, indicative ₹ price, Confirm this pack",
            "  3. Confirm → /verify?code=…  looks up the batch (valid | unverified | invalid | malformed)",
            "  4. Or scan the QR on a physical pack (camera / image / paste ID)",
            "  5. Login as consumer → My checks saves scan history in the browser (per email)",
            "",
            "GROWER PATH (secondary, footer link):",
            "  Photo + herb name + farmer ID + GPS → POST /api/herbs → Flask CNN → hash + QR",
            "  If confidence < 0.70 the batch is stored but marked Not Verified",
            "",
            "AUTH: JWT. New users are role=consumer. Grower demo: farmer@ayurauth.demo",
            "Shopper demo: consumer@ayurauth.demo  /  demo1234@A",
            "No cart or payment — authenticity marketplace, not checkout.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 11 — METHODOLOGY
    # ------------------------------------------------------------------
    {
        "title": "Methodology & Proposed Approach",
        "image": None,
        "bullets": [
            "Step 1 — Shopper discovery:",
            "  Catalog reads GET /api/herbs; default filter = AI-verified only. Confirm pack deep-links to Verify.",
            "",
            "Step 2 — Pack check:",
            "  jsQR decodes camera or uploaded QR. POST /api/verify returns status + herb card (origin, AI %, hash).",
            "",
            "Step 3 — Grower listing (optional):",
            "  Geolocation API + image → Flask CNN. Threshold 0.70. UUID, hash, QR written to local-db.json or MongoDB.",
            "",
            "Step 4 — Account:",
            "  JWT register/login. Consumers → /consumer. Growers → /dashboard.",
            "",
            "Step 5 — Assistant:",
            "  POST /api/chat (rules in chatbotMatch.js) with local React fallback. Information-only; no phone/email.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 12 — CNN MODEL
    # ------------------------------------------------------------------
    {
        "title": "CNN Model — MobileNetV2 (measured on Kaggle leaves)",
        "image": None,
        "bullets": [
            "Weights: herbal_auth_improved_final.h5  |  Flask ml-service  |  input 224×224 RGB /255",
            "Base: ImageNet MobileNetV2, freeze then fine-tune last 30 layers. Softmax → 5 classes.",
            "Classes: Amla, Guava, Hibiscus, Neem, Tulsi  (Kaggle aryashah2k medicinal leaves)",
            "Train: 622 images extracted; metrics on 20% val split, seed 42, n = 122",
            "",
            "Validation scores (sklearn from saved confusion matrix):",
            "  Accuracy 81.15%   |   Macro P/R/F1  79.23% / 78.44% / 78.25%",
            "  Weighted F1 80.83%",
            "  Hibiscus F1 91.30%   Neem 88.46%   Tulsi 80.00%   Guava 79.31%   Amla 52.17%",
            "  Weak class: Amla (7 of 13 predicted as Guava). Train acc ~92% → some overfitting.",
            "  Serve rule: isAuthentic if confidence ≥ 0.70 (AI_VERIFY_THRESHOLD)",
            "  Confusion matrix file: CNN model/files/saved_models/confusion_matrix_final.png",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 13 — AI CHATBOT  (NEW)
    # ------------------------------------------------------------------
    {
        "title": "AI Chatbot — Shopper Assistant",
        "image": None,
        "bullets": [
            "Audience: buyers asking how to scan a pack, read AI scores, or use Shop / My checks.",
            "",
            "Pipeline: message → POST /api/chat → chatbotMatch.js topics → React bubbles",
            "  Fallback: src/lib/chatbot.ts if the API is down. Contact/phone/email requests are refused.",
            "",
            "UI: AIChatbot.tsx floating widget on every page; ChatPage.tsx at /chat (login).",
            "Quick chips: How do I scan a pack?  AI authenticity?  Browse verified herbs?  My checks?",
            "",
            "Topics: scan/verify, AI/CNN, shop listings, grower list-a-batch, hash records, workflow.",
            "Not in scope: placing orders, helpline numbers, medical advice.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 14 — BACKEND  (NEW)
    # ------------------------------------------------------------------
    {
        "title": "Backend — Express, JWT, Local DB (and optional Mongo)",
        "image": None,
        "bullets": [
            "API  http://localhost:5000   (must be running or /shop shows ERR_CONNECTION_REFUSED)",
            "  GET  /api/health   GET/POST /api/herbs   GET /api/herbs/:id   GET .../qr-image",
            "  POST /api/verify   POST /api/auth/register|login   GET /api/auth/me",
            "  POST /api/chat     POST /api/predict (proxy to Flask)",
            "",
            "Auth: JWT. Register default role = consumer. Grower role still allowed.",
            "Storage: USE_LOCAL_DB=true → backend/data/local-db.json (used when Atlas DNS fails).",
            "Images: local /uploads in dev. Flask: FLASK_API_URL=http://127.0.0.1:5001",
            "",
            "Herb record: batchId, herbName, farmerId, imageUrl, origin, lat/lng, isVerified,",
            "  aiConfidence, predictedClass, blockchainHash, qrPayload, uploadedBy",
            "Demo shopper: consumer@ayurauth.demo   Demo grower: farmer@ayurauth.demo",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 15 — WORK DONE
    # ------------------------------------------------------------------
    {
        "title": "Work Done — Live B2C Website",
        "image": None,
        "bullets": [
            "B2C positioning: landing, nav, copy, and default signup are shopper-first.",
            "Routes: / shop, verify, consumer, upload, dashboard, login, signup, chat.",
            "Shop catalog with verified filter, origin, AI score, indicative price, Confirm pack.",
            "Verify: camera + file QR + batch ID + query ?code= from Shop.",
            "JWT consumer vs grower redirects. Demo shopper account seeded in local-db.json.",
            "Flask CNN service + improved MobileNetV2 weights; 5-class labels file.",
            "Measured val. accuracy 81.15%, macro F1 78.25%, confusion matrix saved.",
            "Chatbot rewritten for scan/shop questions; 10-language i18n (English B2C strings).",
            "Local file DB fallback so the site works without MongoDB Atlas.",
            "Not built: payment checkout, live Hyperledger peers, export-compliance stamps.",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 16 — SCREENSHOTS (placeholder)
    # ------------------------------------------------------------------
    {
        "title": "Website — Screens to Demo",
        "image": None,
        "bullets": [
            "Run:  npm run dev  (Vite :5173)   and   npm run backend  (Express :5000)",
            "Optional:  python ml-service/app.py  (CNN :5001)",
            "",
            "Show in this order:",
            "  1. Home /  —  Scan a pack, Browse verified herbs, Sign up (grower link in footer)",
            "  2. /shop  —  verified cards, Confirm this pack → /verify?code=",
            "  3. /verify  —  QR camera / upload / batch ID + status banner",
            "  4. /signup then /consumer  —  My checks scan history",
            "  5. /upload  —  grower list-a-batch (secondary)",
            "  6. Chat widget  —  “How do I scan a pack I bought?”",
            "",
            "Open: http://localhost:5173/",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 17 — CONCLUSION
    # ------------------------------------------------------------------
    {
        "title": "Conclusion & QnA",
        "image": None,
        "bullets": [
            "AyurAuth is a working B2C authenticity app: shoppers scan packs, browse verified listings, and keep check history.",
            "",
            "What ships today:",
            "  • Shopper-first React SPA (Scan, Shop, My checks, JWT consumer accounts)",
            "  • MobileNetV2 CNN — 5 classes, 81.15% validation accuracy, 70% serve threshold",
            "  • Flask inference + Express API + local JSON (or Mongo) + QR (jsQR)",
            "  • Tamper-evident batch hash (simulated ledger, not a live Fabric network)",
            "  • Shopper chatbot + 10-language UI",
            "",
            "Honest next steps:",
            "  • Keep API on :5000 in local demos (USE_LOCAL_DB if Atlas DNS fails)",
            "  • Optional: checkout/payments if you want full e-commerce",
            "  • Optional: real Hyperledger Fabric instead of hash strings",
            "  • Improve Amla class (more images) to lift recall from 46%",
            "",
            "THANK YOU — Questions Welcome!",
        ],
    },

    # ------------------------------------------------------------------
    # SLIDE 18 — REFERENCES
    # ------------------------------------------------------------------
    {
        "title": "References",
        "image": None,
        "bullets": [
            "1. Kunjachan et al. (2023). Comparative Study of CNNs for Leaf Classification in Ayurveda. GEn-CITy 2023.",
            "2. Kaushik et al. (2025). End-to-end traceability in Ayush drug supply chains. Int. J. Ayurveda Research, 6(4), 379–386.",
            "3. Uniyal et al. Leveraging Deep Learning and Blockchain for Herbal Supply Chain Transparency. SRM Institute.",
            "4. Liu et al. (2021). Segmented Encryption: Quality Model for Herbal Medicine using Blockchain. IEEE HEALTHCOM.",
            "5. Rout, S.K. (2025). MedicLedger: Blockchain-Based Architecture for Ayurvedic Authentication. IJCET 16(1).",
            "6. Balaji & Parani (2022). DNA Barcoding of Herbal Powders Reveals Adulteration. Diversity, 14(495).",
            "7. Liu et al. (2021). Traceability for Rare Chinese Herbs using Blockchain. 3rd Int. Conf. AI & Adv. Manufacture.",
            "8. Kaushal et al. (2023). Demystifying Hyperledger Fabric. ICSEIET, 116–120.",
            "9. Reddy et al. (2024). AI & Data Analytics for Pharmaceutical Supply Chain Optimization. IJISAE 12, 934–941.",
            "10. Yik et al. (2021). HerBChain: Blockchain Platform for Herbal Quality Assurance. JTCM 11(6), 598–600.",
        ],
    },
]


OUTPUT_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "AyurAuth_Presentation.pptx")


# =================================================================
#  EXACT DIMENSIONS  (all in EMU via python-pptx helpers)
# =================================================================
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

LR = Cm(1.0)
TM = Cm(3.5)
BM = Cm(2.5)

CW = SLIDE_W - 2 * LR
CH = SLIDE_H - TM - BM

IW = int(CW * 0.70)
IH = int(CH)

BUL_W  = int(CW * 0.30)
BUL_X  = LR + IW
BUL_GAP = Cm(0.3)

TW = CW

WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
FONT  = "Times New Roman"
FOOTER_LABEL = "AyurAuth — Aditya Rajesh Srinivas & Shreya Vikas Jhamb"


# =================================================================
#  HELPERS
# =================================================================

def _set_run_font(run, size_pt, bold=False, color=BLACK,
                  italic=False, alpha_pct=None):
    run.font.name   = FONT
    run.font.size   = Pt(size_pt)
    run.font.bold   = bold
    run.font.italic = italic

    if alpha_pct is not None:
        alpha_val = int(alpha_pct * 1000)
        rPr = run._r.get_or_add_rPr()
        for old in rPr.findall(qn('a:solidFill')):
            rPr.remove(old)
        solid  = etree.SubElement(rPr, qn('a:solidFill'))
        srgb   = etree.SubElement(solid, qn('a:srgbClr'))
        hex_c  = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
        srgb.set('val', hex_c)
        alpha_e = etree.SubElement(srgb, qn('a:alpha'))
        alpha_e.set('val', str(alpha_val))
    else:
        run.font.color.rgb = color


def _add_textbox(slide, left, top, width, height):
    shape = slide.shapes.add_textbox(int(left), int(top),
                                     int(width), int(height))
    tf = shape.text_frame
    tf.word_wrap     = True
    tf.margin_left   = Pt(0)
    tf.margin_right  = Pt(0)
    tf.margin_top    = Pt(0)
    tf.margin_bottom = Pt(0)
    return shape, tf


def _white_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def _add_background_image(slide, bg_path="background.jpg"):
    if os.path.exists(bg_path):
        slide.shapes.add_picture(bg_path, 0, 0, SLIDE_W, SLIDE_H)


# =================================================================
#  TITLE  —  36 pt bold, centred in top 3.5 cm
# =================================================================

def _draw_title(slide, title_text):
    shape, tf = _add_textbox(slide, LR, 0, TW, TM)
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title_text
    _set_run_font(run, 36, bold=True, color=BLACK)


# =================================================================
#  IMAGE  —  auto aspect ratio, never exceeds IW x IH
# =================================================================

def _get_fitted_image_size(image_path):
    with PILImage.open(image_path) as img:
        w_px, h_px = img.size
    aspect = w_px / h_px
    fit_w  = IW
    fit_h  = int(fit_w / aspect)
    if fit_h > IH:
        fit_h = IH
        fit_w = int(fit_h * aspect)
    fit_w = min(fit_w, IW)
    fit_h = min(fit_h, IH)
    return fit_w, fit_h


def _draw_image(slide, image_path):
    if not image_path or not os.path.exists(image_path):
        return
    pic_w, pic_h = _get_fitted_image_size(image_path)
    pic_left = int(LR + (IW - pic_w) / 2)
    pic_top  = int(TM + (IH - pic_h) / 2)
    slide.shapes.add_picture(image_path, pic_left, pic_top, pic_w, pic_h)


# =================================================================
#  BULLETS  —  18 pt, right column (or full width if no image)
# =================================================================

def _draw_bullets(slide, bullets, image_present):
    col_left  = int(BUL_X + BUL_GAP) if image_present else int(LR)
    col_width = int(BUL_W - BUL_GAP) if image_present else int(CW)

    _, tf = _add_textbox(slide, col_left, TM, col_width, CH)

    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment    = PP_ALIGN.LEFT
        p.space_before = Pt(2)
        p.space_after  = Pt(2)

        run = p.add_run()
        run.text = bullet if bullet else ""
        # Slightly smaller font (18 pt) so long bullet lists fit neatly
        _set_run_font(run, 18)


# =================================================================
#  FOOTER  —  10 pt, 50 % opacity, centred in bottom margin
# =================================================================

def _draw_footer(slide, slide_number, total_slides):
    footer_top = SLIDE_H - Cm(1.5)
    _, tf = _add_textbox(slide, LR, footer_top, CW, Cm(0.9))
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.BOTTOM
    tf.word_wrap = False

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    run = p.add_run()
    run.text = f"Slide\u00a0{slide_number}\u2003\u2003{FOOTER_LABEL}"
    _set_run_font(run, 10, alpha_pct=50)


# =================================================================
#  BUILD PRESENTATION
# =================================================================

def build(slides_data, output_file=OUTPUT_FILE):
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    total = len(slides_data)

    for idx, data in enumerate(slides_data, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank

        _white_background(slide)
        _add_background_image(slide)

        _draw_title(slide, data.get("title", "Untitled"))

        img_path      = data.get("image")
        image_present = bool(img_path and os.path.exists(str(img_path)))

        _draw_image(slide, img_path)
        _draw_bullets(slide, data.get("bullets", []), image_present)
        _draw_footer(slide, idx, total)

        print(f"  Slide {idx}/{total}  --  {data.get('title','')}")

    prs.save(output_file)
    print(f"\nSaved: {output_file}")


# =================================================================
if __name__ == "__main__":
    build(SLIDES)
